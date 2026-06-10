import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from io import BytesIO
import warnings
warnings.filterwarnings("ignore")

# ── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Profit Trends Dashboard",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Styling ───────────────────────────────────────────────────────────────────
st.markdown("""
<style>
  [data-testid="stMetricValue"] { font-size: 1.6rem; font-weight: 700; }
  .kpi-row { gap: 1rem; }
  .stTabs [data-baseweb="tab"] { font-size: 0.9rem; font-weight: 600; }
  .tag-grow  { background:#d1fae5; color:#065f46; padding:2px 8px; border-radius:4px; font-size:.78rem; }
  .tag-dec   { background:#fee2e2; color:#991b1b; padding:2px 8px; border-radius:4px; font-size:.78rem; }
  .tag-watch { background:#fef9c3; color:#713f12; padding:2px 8px; border-radius:4px; font-size:.78rem; }
  .tag-high  { background:#dbeafe; color:#1e40af; padding:2px 8px; border-radius:4px; font-size:.78rem; }
</style>
""", unsafe_allow_html=True)

COLORS = {
    "primary": "#1e40af",
    "success": "#059669",
    "danger":  "#dc2626",
    "warning": "#d97706",
    "neutral": "#6b7280",
}

# ── Data loading ──────────────────────────────────────────────────────────────
@st.cache_data
def load_data(path: str) -> pd.DataFrame:
    df = pd.read_excel(path)
    df["Booked At (UTC)"] = pd.to_datetime(df["Booked At (UTC)"], utc=True)
    df["Date"]   = df["Booked At (UTC)"].dt.date
    df["Month"]  = df["Booked At (UTC)"].dt.to_period("M").dt.to_timestamp()
    df["Week"]   = df["Booked At (UTC)"].dt.to_period("W").dt.to_timestamp()
    df["MonthLabel"] = df["Booked At (UTC)"].dt.strftime("%b %Y")
    df["Revenue"]     = pd.to_numeric(df["Revenue"],     errors="coerce").fillna(0)
    df["Act. Costs"]  = pd.to_numeric(df["Act. Costs"],  errors="coerce").fillna(0)
    df["Est. Costs"]  = pd.to_numeric(df["Est. Costs"],  errors="coerce").fillna(0)
    df["Est. Profit"] = pd.to_numeric(df["Est. Profit"], errors="coerce").fillna(0)
    df["Fin. Profit"] = pd.to_numeric(df["Fin. Profit"], errors="coerce").fillna(0)
    df["Location"]    = df["Location"].fillna("Unknown")
    return df

@st.cache_data
def build_customer_summary(df: pd.DataFrame) -> pd.DataFrame:
    grp = df.groupby("Customer").agg(
        Revenue=("Revenue", "sum"),
        Act_Costs=("Act. Costs", "sum"),
        Est_Profit=("Est. Profit", "sum"),
        Fin_Profit=("Fin. Profit", "sum"),
        Shipments=("BOL #", "count"),
    ).reset_index()
    grp["Avg_Profit"]  = grp["Fin_Profit"] / grp["Shipments"]
    grp["Margin"]      = np.where(grp["Revenue"] != 0, grp["Fin_Profit"] / grp["Revenue"] * 100, 0)
    return grp

@st.cache_data
def build_monthly_customer(df: pd.DataFrame) -> pd.DataFrame:
    return df.groupby(["Customer", "Month"]).agg(
        Revenue=("Revenue", "sum"),
        Fin_Profit=("Fin. Profit", "sum"),
        Shipments=("BOL #", "count"),
    ).reset_index().sort_values(["Customer", "Month"])

@st.cache_data
def flag_customers(mc: pd.DataFrame, summary: pd.DataFrame) -> pd.DataFrame:
    customers = summary["Customer"].unique()
    flags = []
    for cust in customers:
        cdf = mc[mc["Customer"] == cust].sort_values("Month")
        if len(cdf) < 2:
            flags.append({"Customer": cust, "Flag": "Insufficient data", "MoM_Change": None})
            continue
        latest = cdf.iloc[-1]["Fin_Profit"]
        prev   = cdf.iloc[-2]["Fin_Profit"]
        n3     = cdf.tail(3)["Fin_Profit"].mean() if len(cdf) >= 3 else prev
        mom    = (latest - prev) / abs(prev) if prev != 0 else 0
        vs3m   = (latest - n3)   / abs(n3)   if n3   != 0 else 0
        total  = summary.loc[summary["Customer"] == cust, "Fin_Profit"].values[0]
        margin_now  = (cdf.iloc[-1]["Fin_Profit"] / cdf.iloc[-1]["Revenue"] * 100) if cdf.iloc[-1]["Revenue"] != 0 else 0
        margin_prev = (cdf.iloc[-2]["Fin_Profit"] / cdf.iloc[-2]["Revenue"] * 100) if cdf.iloc[-2]["Revenue"] != 0 else 0
        ship_now    = cdf.iloc[-1]["Shipments"]
        ship_prev   = cdf.iloc[-2]["Shipments"]

        if total > 500:
            flag = "High Value"
        elif mom >= 0.15 or vs3m >= 0.15:
            flag = "Growing"
        elif mom <= -0.15 or vs3m <= -0.15:
            flag = "Declining"
        elif latest > 0 and (mom < 0 or margin_now < margin_prev or ship_now < ship_prev):
            flag = "Watchlist"
        else:
            flag = "Stable"

        flags.append({
            "Customer":     cust,
            "Flag":         flag,
            "MoM_Change":   mom,
            "vs3M_Change":  vs3m,
            "Latest_Profit": latest,
            "Prev_Profit":   prev,
            "Margin_Now":    margin_now,
            "Margin_Prev":   margin_prev,
            "Ship_Now":      ship_now,
            "Ship_Prev":     ship_prev,
        })
    flags_df = pd.DataFrame(flags)
    return summary.merge(flags_df, on="Customer", how="left")

# ── Report generation ─────────────────────────────────────────────────────────
def generate_report(flagged: pd.DataFrame) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_style(cell, bg="1e40af", fg="FFFFFF", sz=11, bold=True):
        cell.font = Font(bold=bold, color=fg, size=sz)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    def cell_style(cell, align="left"):
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border = border

    def write_section(ws, title, cols, rows, start_row, title_color="1e40af"):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(cols))
        tc = ws.cell(start_row, 1, title)
        tc.font = Font(bold=True, size=12, color="FFFFFF")
        tc.fill = PatternFill("solid", fgColor=title_color)
        tc.alignment = Alignment(horizontal="left", vertical="center")
        start_row += 1
        for ci, col in enumerate(cols, 1):
            c = ws.cell(start_row, ci, col)
            hdr_style(c, bg="374151", sz=10)
        start_row += 1
        for row in rows:
            for ci, val in enumerate(row, 1):
                c = ws.cell(start_row, ci, val)
                cell_style(c, "right" if ci > 1 else "left")
            start_row += 1
        return start_row + 2

    # ── Sheet 1: Executive Summary ────────────────────────────────────────────
    ws = wb.active
    ws.title = "Executive Summary"
    ws.column_dimensions["A"].width = 32
    for col in ["B","C","D","E"]:
        ws.column_dimensions[col].width = 18

    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = "2026 Customer Profit Trend Report — Executive Summary"
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1e40af")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    metrics = [
        ("Total Revenue",       f"${flagged['Revenue'].sum():,.0f}"),
        ("Total Fin. Profit",   f"${flagged['Fin_Profit'].sum():,.0f}"),
        ("Overall Profit Margin", f"{flagged['Fin_Profit'].sum()/flagged['Revenue'].sum()*100:.1f}%"),
        ("Total Shipments",     f"{flagged['Shipments'].sum():,}"),
        ("Active Customers",    str(len(flagged))),
        ("Growing Customers",   str((flagged['Flag']=='Growing').sum())),
        ("Declining Customers", str((flagged['Flag']=='Declining').sum())),
        ("High-Value Customers",str((flagged['Flag']=='High Value').sum())),
        ("Watchlist Customers", str((flagged['Flag']=='Watchlist').sum())),
    ]
    for i, (lbl, val) in enumerate(metrics, 3):
        c1 = ws.cell(i, 1, lbl); c1.font = Font(bold=True, size=10); c1.border = border
        c2 = ws.cell(i, 2, val); c2.font = Font(size=10); c2.border = border
        c2.alignment = Alignment(horizontal="right")

    flag_colors = {"Growing":"065f46","Declining":"991b1b","High Value":"1e3a5f","Watchlist":"713f12","Stable":"374151"}
    row = 14
    ws.cell(row, 1, "KEY FINDINGS").font = Font(bold=True, size=11, color="1e40af")
    row += 1
    findings = [
        f"• {(flagged['Flag']=='Growing').sum()} customers show strong growth (≥15% MoM or vs 3M avg).",
        f"• {(flagged['Flag']=='Declining').sum()} customers are declining — require immediate attention.",
        f"• {(flagged['Flag']=='High Value').sum()} high-value accounts (>$500 total profit) identified.",
        f"• {(flagged['Flag']=='Watchlist').sum()} accounts on watchlist with shrinking margin or volume.",
        f"• Overall profit margin: {flagged['Fin_Profit'].sum()/flagged['Revenue'].sum()*100:.1f}%.",
    ]
    for f in findings:
        c = ws.cell(row, 1, f); c.font = Font(size=10); c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 18
        row += 1

    # ── Sheet 2: All Customers ────────────────────────────────────────────────
    ws2 = wb.create_sheet("All Customers")
    cols2 = ["Customer","Flag","Total Revenue","Total Fin. Profit","Shipments","Avg Profit/Ship","Margin %","MoM Change"]
    widths2 = [36,14,16,18,12,18,12,14]
    for i,(c,w) in enumerate(zip(cols2,widths2),1):
        ws2.column_dimensions[get_column_letter(i)].width = w
        cell = ws2.cell(1,i,c); hdr_style(cell)

    for ri, row in flagged.sort_values("Fin_Profit",ascending=False).iterrows():
        r = ws2.max_row + 1
        vals = [
            row["Customer"], row["Flag"],
            round(row["Revenue"],2), round(row["Fin_Profit"],2),
            int(row["Shipments"]), round(row.get("Avg_Profit",0),2),
            round(row.get("Margin",0),1),
            f"{row['MoM_Change']*100:.1f}%" if pd.notna(row.get("MoM_Change")) else "N/A",
        ]
        for ci,v in enumerate(vals,1):
            c = ws2.cell(r,ci,v); cell_style(c, "right" if ci>1 else "left")
        # colour flag cell
        fc = ws2.cell(r,2)
        bg = flag_colors.get(row["Flag"],"374151")
        fg = "FFFFFF"
        fc.fill = PatternFill("solid", fgColor=bg)
        fc.font = Font(color=fg, size=9, bold=True)

    # ── Sheet 3: Flagged Segments ─────────────────────────────────────────────
    ws3 = wb.create_sheet("Flagged Accounts")
    ws3.column_dimensions["A"].width = 36
    for col in ["B","C","D","E","F"]:
        ws3.column_dimensions[col].width = 18

    seg_row = 1
    for flag_val, color in [("Growing","065f46"),("Declining","991b1b"),("High Value","1e3a5f"),("Watchlist","713f12")]:
        sub = flagged[flagged["Flag"]==flag_val].sort_values("Fin_Profit",ascending=False)
        if sub.empty: continue
        rows_data = [
            (r["Customer"], f"${r['Revenue']:,.0f}", f"${r['Fin_Profit']:,.0f}",
             f"{r.get('Margin',0):.1f}%",
             f"{r['MoM_Change']*100:+.1f}%" if pd.notna(r.get("MoM_Change")) else "N/A",
             int(r["Shipments"]))
            for _,r in sub.iterrows()
        ]
        seg_row = write_section(
            ws3,
            f"{'🟢' if flag_val=='Growing' else '🔴' if flag_val=='Declining' else '🔵' if flag_val=='High Value' else '🟡'} {flag_val} Accounts",
            ["Customer","Revenue","Fin. Profit","Margin","MoM Change","Shipments"],
            rows_data, seg_row, title_color=color,
        )

    # ── Sheet 4: Recommendations ──────────────────────────────────────────────
    ws4 = wb.create_sheet("Recommendations")
    ws4.column_dimensions["A"].width = 36
    ws4.column_dimensions["B"].width = 18
    ws4.column_dimensions["C"].width = 60
    ws4.merge_cells("A1:C1")
    h = ws4["A1"]
    h.value = "Recommended Follow-Up Actions"
    h.font = Font(bold=True, size=13, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor="1e40af")
    h.alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 28
    for ci,lbl in enumerate(["Customer","Flag","Recommended Action"],1):
        c = ws4.cell(2,ci,lbl); hdr_style(c)

    recommendations = {
        "Growing":    "Prioritize for upselling; negotiate volume commitment or preferred rates. Schedule quarterly business review.",
        "Declining":  "Urgent outreach required. Identify root cause (pricing? competition? demand drop?). Offer retention incentive.",
        "High Value": "Maintain strong service levels. Assign dedicated account manager. Explore cross-sell opportunities.",
        "Watchlist":  "Monitor closely for next 30 days. Initiate proactive check-in call. Review contract terms and pricing alignment.",
        "Stable":     "Maintain current service. Include in next newsletter or market update.",
    }
    for _,row in flagged.sort_values("Fin_Profit",ascending=False).iterrows():
        r = ws4.max_row + 1
        for ci,v in enumerate([row["Customer"], row["Flag"], recommendations.get(row["Flag"],"Review account.")],1):
            c = ws4.cell(r,ci,v)
            cell_style(c, "left")
            ws4.row_dimensions[r].height = 20
        fc = ws4.cell(r,2)
        fc.fill = PatternFill("solid", fgColor=flag_colors.get(row["Flag"],"374151"))
        fc.font = Font(color="FFFFFF", size=9, bold=True)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

# ── Load data ─────────────────────────────────────────────────────────────────
DATA_PATH = "2026_Profit_Report.xlsx"

try:
    df_raw = load_data(DATA_PATH)
except FileNotFoundError:
    st.error("⚠️  Could not find **2026_Profit_Report.xlsx**. Place it in the same folder as app.py and refresh.")
    st.stop()

# ── Sidebar filters ───────────────────────────────────────────────────────────
st.sidebar.image("https://img.icons8.com/fluency/48/combo-chart.png", width=40)
st.sidebar.title("Filters")

min_d = df_raw["Month"].min().date()
max_d = df_raw["Month"].max().date()
date_range = st.sidebar.date_input("Date range", value=(min_d, max_d), min_value=min_d, max_value=max_d)
if len(date_range) == 2:
    d_start, d_end = pd.Timestamp(date_range[0]), pd.Timestamp(date_range[1])
else:
    d_start, d_end = pd.Timestamp(min_d), pd.Timestamp(max_d)

customers_all = sorted(df_raw["Customer"].unique())
sel_customers = st.sidebar.multiselect("Customer", customers_all, placeholder="All customers")

types_all = sorted(df_raw["Type"].dropna().unique())
sel_types = st.sidebar.multiselect("Shipment Type", types_all, placeholder="All types")

pickup_all = sorted(df_raw["Pickup Country"].dropna().unique())
sel_pickup = st.sidebar.multiselect("Pickup Country", pickup_all, placeholder="All countries")

delivery_all = sorted(df_raw["Delivery Country"].dropna().unique())
sel_delivery = st.sidebar.multiselect("Delivery Country", delivery_all, placeholder="All countries")

locs_all = sorted(df_raw["Location"].dropna().astype(str).unique())
sel_locs = st.sidebar.multiselect("Location", locs_all, placeholder="All locations")

# Apply filters
df = df_raw[
    (df_raw["Month"] >= d_start) &
    (df_raw["Month"] <= d_end)
]
if sel_customers: df = df[df["Customer"].isin(sel_customers)]
if sel_types:     df = df[df["Type"].isin(sel_types)]
if sel_pickup:    df = df[df["Pickup Country"].isin(sel_pickup)]
if sel_delivery:  df = df[df["Delivery Country"].isin(sel_delivery)]
if sel_locs:      df = df[df["Location"].isin(sel_locs)]

summary    = build_customer_summary(df)
monthly_c  = build_monthly_customer(df)
flagged    = flag_customers(monthly_c, summary)

# ── Header ─────────────────────────────────────────────────────────────────────
st.markdown("## 📊 Customer Profit Trends Dashboard")
st.caption(f"Period: {df['Month'].min().strftime('%b %Y')} – {df['Month'].max().strftime('%b %Y')}  •  {len(df):,} shipments  •  {df['Customer'].nunique()} customers")

# ── KPI Cards ─────────────────────────────────────────────────────────────────
k1,k2,k3,k4,k5 = st.columns(5)
total_rev    = df["Revenue"].sum()
total_profit = df["Fin. Profit"].sum()
margin_pct   = total_profit / total_rev * 100 if total_rev else 0
k1.metric("Total Revenue",       f"${total_rev:,.0f}")
k2.metric("Total Fin. Profit",   f"${total_profit:,.0f}")
k3.metric("Profit Margin",       f"{margin_pct:.1f}%")
k4.metric("Total Shipments",     f"{len(df):,}")
k5.metric("Active Customers",    str(df["Customer"].nunique()))

st.divider()

# ── Tabs ───────────────────────────────────────────────────────────────────────
tab1, tab2, tab3, tab4 = st.tabs(["📈 Trends", "👥 Customers", "🚦 Trend Flags", "📄 Report"])

# ──────────────────────────── TAB 1: TRENDS ──────────────────────────────────
with tab1:
    monthly = df.groupby("Month").agg(
        Revenue=("Revenue","sum"),
        Fin_Profit=("Fin. Profit","sum"),
        Shipments=("BOL #","count"),
    ).reset_index()
    monthly["Margin"] = monthly["Fin_Profit"] / monthly["Revenue"] * 100

    col1, col2 = st.columns(2)
    with col1:
        st.markdown("#### Monthly Profit Trend")
        fig = px.line(monthly, x="Month", y="Fin_Profit", markers=True,
                      color_discrete_sequence=[COLORS["primary"]])
        fig.update_layout(xaxis_title="", yaxis_title="Fin. Profit ($)", height=300,
                          plot_bgcolor="white", paper_bgcolor="white",
                          yaxis=dict(gridcolor="#f3f4f6"))
        fig.update_traces(line_width=2.5)
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        st.markdown("#### Revenue vs Profit")
        fig2 = go.Figure()
        fig2.add_trace(go.Bar(x=monthly["Month"], y=monthly["Revenue"], name="Revenue",
                               marker_color="#bfdbfe"))
        fig2.add_trace(go.Bar(x=monthly["Month"], y=monthly["Fin_Profit"], name="Fin. Profit",
                               marker_color=COLORS["primary"]))
        fig2.update_layout(barmode="overlay", xaxis_title="", yaxis_title="$",
                           height=300, plot_bgcolor="white", paper_bgcolor="white",
                           yaxis=dict(gridcolor="#f3f4f6"), legend=dict(orientation="h", y=1.1))
        st.plotly_chart(fig2, use_container_width=True)

    col3, col4 = st.columns(2)
    with col3:
        st.markdown("#### Shipment Count by Month")
        fig3 = px.bar(monthly, x="Month", y="Shipments",
                      color_discrete_sequence=[COLORS["neutral"]])
        fig3.update_layout(xaxis_title="", yaxis_title="Shipments", height=280,
                           plot_bgcolor="white", paper_bgcolor="white",
                           yaxis=dict(gridcolor="#f3f4f6"))
        st.plotly_chart(fig3, use_container_width=True)

    with col4:
        st.markdown("#### Profit Margin % by Month")
        fig4 = px.line(monthly, x="Month", y="Margin", markers=True,
                       color_discrete_sequence=[COLORS["success"]])
        fig4.update_layout(xaxis_title="", yaxis_title="Margin (%)", height=280,
                           plot_bgcolor="white", paper_bgcolor="white",
                           yaxis=dict(gridcolor="#f3f4f6"))
        fig4.add_hline(y=monthly["Margin"].mean(), line_dash="dot", line_color="gray",
                       annotation_text="Avg", annotation_position="right")
        st.plotly_chart(fig4, use_container_width=True)

# ──────────────────────────── TAB 2: CUSTOMERS ───────────────────────────────
with tab2:
    c1, c2 = st.columns(2)

    with c1:
        st.markdown("#### Top 20 Customers by Fin. Profit")
        top20 = summary.nlargest(20, "Fin_Profit")
        fig5 = px.bar(top20.sort_values("Fin_Profit"), x="Fin_Profit", y="Customer",
                      orientation="h", color="Fin_Profit",
                      color_continuous_scale=["#bfdbfe","#1e40af"])
        fig5.update_layout(height=520, xaxis_title="Fin. Profit ($)", yaxis_title="",
                           coloraxis_showscale=False, plot_bgcolor="white", paper_bgcolor="white")
        st.plotly_chart(fig5, use_container_width=True)

    with c2:
        st.markdown("#### Profit Margin by Customer (Top 30)")
        top30 = summary.nlargest(30, "Fin_Profit")
        fig6 = px.bar(top30.sort_values("Margin"), x="Margin", y="Customer",
                      orientation="h", color="Margin",
                      color_continuous_scale=["#fca5a5","#86efac"])
        fig6.update_layout(height=520, xaxis_title="Margin (%)", yaxis_title="",
                           coloraxis_showscale=False, plot_bgcolor="white", paper_bgcolor="white")
        st.plotly_chart(fig6, use_container_width=True)

    st.markdown("#### Bottom / Declining Customers (Lowest Fin. Profit)")
    bottom15 = summary.nsmallest(15, "Fin_Profit")[["Customer","Revenue","Fin_Profit","Shipments","Margin","Avg_Profit"]]
    bottom15.columns = ["Customer","Revenue ($)","Fin. Profit ($)","Shipments","Margin (%)","Avg Profit/Ship ($)"]
    for col in ["Revenue ($)","Fin. Profit ($)","Avg Profit/Ship ($)"]:
        bottom15[col] = bottom15[col].map("${:,.2f}".format)
    bottom15["Margin (%)"] = bottom15["Margin (%)"].map("{:.1f}%".format)
    st.dataframe(bottom15.reset_index(drop=True), use_container_width=True, height=320)

# ──────────────────────────── TAB 3: TREND FLAGS ─────────────────────────────
with tab3:
    # MoM growth / decline
    mc_sorted = monthly_c.sort_values(["Customer","Month"])
    mc_sorted["Prev_Profit"] = mc_sorted.groupby("Customer")["Fin_Profit"].shift(1)
    mc_sorted["MoM"] = (mc_sorted["Fin_Profit"] - mc_sorted["Prev_Profit"]) / mc_sorted["Prev_Profit"].abs()
    latest_month = mc_sorted["Month"].max()
    latest_mc    = mc_sorted[mc_sorted["Month"] == latest_month].dropna(subset=["MoM"])

    col_g, col_d = st.columns(2)

    with col_g:
        st.markdown("#### 🟢 Biggest MoM Profit Growth")
        growers = latest_mc.nlargest(10, "MoM")[["Customer","Fin_Profit","Prev_Profit","MoM"]]
        growers.columns = ["Customer","This Month ($)","Prev Month ($)","MoM Change"]
        growers["This Month ($)"]  = growers["This Month ($)"].map("${:,.0f}".format)
        growers["Prev Month ($)"]  = growers["Prev Month ($)"].map("${:,.0f}".format)
        growers["MoM Change"]      = growers["MoM Change"].map("{:+.1%}".format)
        st.dataframe(growers.reset_index(drop=True), use_container_width=True)

    with col_d:
        st.markdown("#### 🔴 Biggest MoM Profit Decline")
        decliners = latest_mc.nsmallest(10, "MoM")[["Customer","Fin_Profit","Prev_Profit","MoM"]]
        decliners.columns = ["Customer","This Month ($)","Prev Month ($)","MoM Change"]
        decliners["This Month ($)"] = decliners["This Month ($)"].map("${:,.0f}".format)
        decliners["Prev Month ($)"] = decliners["Prev Month ($)"].map("${:,.0f}".format)
        decliners["MoM Change"]     = decliners["MoM Change"].map("{:+.1%}".format)
        st.dataframe(decliners.reset_index(drop=True), use_container_width=True)

    st.divider()
    st.markdown("#### Customer Trend Flags — Full Overview")
    flag_order = ["High Value","Growing","Watchlist","Declining","Stable","Insufficient data"]
    flag_colors_map = {
        "Growing":"🟢","Declining":"🔴","High Value":"🔵","Watchlist":"🟡","Stable":"⚪","Insufficient data":"⚫"
    }
    display_flagged = flagged.copy()
    display_flagged["Flag"] = display_flagged["Flag"].map(lambda x: f"{flag_colors_map.get(x,'')} {x}")
    display_flagged = display_flagged.sort_values(["Fin_Profit"], ascending=False)
    cols_show = ["Customer","Flag","Revenue","Fin_Profit","Shipments","Margin","MoM_Change"]
    df_show = display_flagged[cols_show].copy()
    df_show.columns = ["Customer","Flag","Revenue ($)","Fin. Profit ($)","Shipments","Margin (%)","MoM Change"]
    df_show["Revenue ($)"]      = df_show["Revenue ($)"].map("${:,.0f}".format)
    df_show["Fin. Profit ($)"]  = df_show["Fin. Profit ($)"].map("${:,.0f}".format)
    df_show["Margin (%)"]       = df_show["Margin (%)"].map("{:.1f}%".format)
    df_show["MoM Change"]       = df_show["MoM Change"].apply(
        lambda x: f"{x*100:+.1f}%" if pd.notna(x) else "N/A"
    )
    st.dataframe(df_show.reset_index(drop=True), use_container_width=True, height=480)

    # Scatter: Revenue vs Profit coloured by flag
    st.markdown("#### Revenue vs Profit by Customer (colour = flag)")
    scatter_df = flagged[flagged["Revenue"] > 0].copy()
    scatter_df["Flag_clean"] = scatter_df["Flag"]
    fig_sc = px.scatter(
        scatter_df, x="Revenue", y="Fin_Profit", color="Flag_clean",
        hover_name="Customer", size="Shipments",
        color_discrete_map={
            "Growing":"#059669","Declining":"#dc2626",
            "High Value":"#1e40af","Watchlist":"#d97706","Stable":"#9ca3af",
        },
        labels={"Revenue":"Revenue ($)","Fin_Profit":"Fin. Profit ($)","Flag_clean":"Flag"},
    )
    fig_sc.update_layout(height=400, plot_bgcolor="white", paper_bgcolor="white")
    st.plotly_chart(fig_sc, use_container_width=True)

# ──────────────────────────── TAB 4: REPORT ──────────────────────────────────
with tab4:
    st.markdown("#### 📄 Downloadable Account Report")
    st.markdown("""
    The report includes:
    - **Executive Summary** with KPIs and key findings
    - **All Customers** sheet with flags and metrics
    - **Flagged Accounts** by category (Growing, Declining, High Value, Watchlist)
    - **Recommendations** with specific follow-up actions per customer
    """)

    if st.button("Generate Report", type="primary"):
        with st.spinner("Building report…"):
            report_bytes = generate_report(flagged)
        st.download_button(
            label="⬇️ Download Excel Report",
            data=report_bytes,
            file_name="Customer_Profit_Report_2026.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        st.success("Report ready — click above to download.")

    st.divider()
    st.markdown("#### Flag Summary")
    flag_summary = flagged.groupby("Flag").agg(
        Customers=("Customer","count"),
        Total_Profit=("Fin_Profit","sum"),
        Avg_Margin=("Margin","mean"),
    ).reset_index().sort_values("Total_Profit", ascending=False)
    flag_summary.columns = ["Flag","# Customers","Total Fin. Profit ($)","Avg Margin (%)"]
    flag_summary["Total Fin. Profit ($)"] = flag_summary["Total Fin. Profit ($)"].map("${:,.0f}".format)
    flag_summary["Avg Margin (%)"] = flag_summary["Avg Margin (%)"].map("{:.1f}%".format)
    st.dataframe(flag_summary.reset_index(drop=True), use_container_width=True)
